"""Document text extraction (US-009).

One function per supported format (US-008's acceptance criteria: PDF,
DOCX, TXT, Markdown, CSV). PDF extraction tries native text first
(fast, works for the overwhelming majority of real documents) and
only pays the OCR cost for pages that come back empty — i.e. scanned
or image-only pages.

Every extractor returns a list of ``(page_number, text)`` pairs rather
than one flat string: PDFs get one entry per page (1-indexed) so
Epic 4's source attribution (US-019) can report which page a chunk
came from; formats with no native pagination (DOCX/TXT/MD/CSV) return
a single entry with ``page_number=None``.
"""

from __future__ import annotations

import csv
import logging
from pathlib import Path

logger = logging.getLogger(__name__)

# A native-text page with fewer than this many stripped characters is
# treated as "probably scanned" and rasterised for OCR instead. Real
# text pages are almost never this sparse; scanned pages with no
# extractable text layer come back completely empty (0 chars).
_MIN_NATIVE_TEXT_CHARS = 30

PageText = tuple[int | None, str]


def extract_pages(path: Path, file_type: str) -> list[PageText]:
    if file_type == "pdf":
        return _extract_pdf(path)
    if file_type == "docx":
        return [(None, _extract_docx(path))]
    if file_type in ("txt", "md"):
        return [(None, _extract_plain_text(path))]
    if file_type == "csv":
        return [(None, _extract_csv(path))]
    raise ValueError(f"Unsupported file_type '{file_type}'")


def _extract_pdf(path: Path) -> list[PageText]:
    import fitz  # pymupdf

    from app.indexing.ocr import ocr_image

    pages: list[PageText] = []
    doc = fitz.open(path)
    try:
        for page in doc:
            text = page.get_text()
            if len(text.strip()) < _MIN_NATIVE_TEXT_CHARS:
                logger.info(
                    "page %d of %s has little/no native text; falling back to OCR",
                    page.number,
                    path.name,
                )
                pixmap = page.get_pixmap(dpi=200)
                text = ocr_image(pixmap.tobytes("png"))
            # pymupdf's page.number is 0-indexed; report 1-indexed page
            # numbers since that's what a human (or a "page 17" source
            # citation) expects.
            pages.append((page.number + 1, text))
    finally:
        doc.close()
    return pages


def _extract_docx(path: Path) -> str:
    from docx import Document as DocxDocument

    doc = DocxDocument(path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_plain_text(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")


def _extract_csv(path: Path) -> str:
    """Serialise each row as ``col: value, col: value`` lines rather
    than raw comma-separated text — meaningfully more useful for
    chunking/embedding than unlabelled CSV soup.
    """
    lines: list[str] = []
    with path.open(newline="", encoding="utf-8", errors="replace") as fh:
        reader = csv.DictReader(fh)
        for row in reader:
            lines.append(", ".join(f"{key}: {value}" for key, value in row.items()))
    return "\n".join(lines)
