"""Tests for document text extraction (US-009) and page tracking
(Epic 4 / US-019 source attribution).

PDF/DOCX/TXT/MD/CSV are all exercised against small generated
fixtures with the real parsing libraries — fast and deterministic.
The OCR fallback path (real PaddleOCR) is intentionally NOT exercised
here; see the plan's manual end-to-end verification for that instead
(downloading/running the real model is too slow for a unit suite).
"""

from __future__ import annotations

import csv

import pytest

from app.indexing.parsers import extract_pages


def test_extract_txt_has_no_page_number(tmp_path):
    path = tmp_path / "notes.txt"
    path.write_text("Hello knowledge base.\nSecond line.", encoding="utf-8")
    pages = extract_pages(path, "txt")
    assert pages == [(None, "Hello knowledge base.\nSecond line.")]


def test_extract_markdown_has_no_page_number(tmp_path):
    path = tmp_path / "guide.md"
    content = "# Title\n\nSome body text."
    path.write_text(content, encoding="utf-8")
    assert extract_pages(path, "md") == [(None, content)]


def test_extract_csv(tmp_path):
    path = tmp_path / "data.csv"
    with path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow(["name", "status"])
        writer.writerow(["Alice", "active"])
        writer.writerow(["Bob", "inactive"])

    pages = extract_pages(path, "csv")
    assert len(pages) == 1
    page_number, text = pages[0]
    assert page_number is None
    lines = text.splitlines()
    assert lines[0] == "name: Alice, status: active"
    assert lines[1] == "name: Bob, status: inactive"


def test_extract_docx_has_no_page_number(tmp_path):
    from docx import Document as DocxDocument

    path = tmp_path / "report.docx"
    doc = DocxDocument()
    doc.add_paragraph("First paragraph.")
    doc.add_paragraph("Second paragraph.")
    doc.save(path)

    pages = extract_pages(path, "docx")
    assert len(pages) == 1
    page_number, text = pages[0]
    assert page_number is None
    assert "First paragraph." in text
    assert "Second paragraph." in text


def test_extract_pdf_returns_one_entry_per_page_1_indexed(tmp_path):
    import fitz  # pymupdf

    path = tmp_path / "doc.pdf"
    doc = fitz.open()
    page1 = doc.new_page()
    page1.insert_text((72, 72), "First page content.")
    page2 = doc.new_page()
    page2.insert_text((72, 72), "Second page content.")
    doc.save(path)
    doc.close()

    pages = extract_pages(path, "pdf")
    assert len(pages) == 2
    assert pages[0][0] == 1
    assert "First page content." in pages[0][1]
    assert pages[1][0] == 2
    assert "Second page content." in pages[1][1]


def test_unsupported_file_type_raises(tmp_path):
    path = tmp_path / "a.xyz"
    path.write_text("data", encoding="utf-8")
    with pytest.raises(ValueError):
        extract_pages(path, "xyz")
